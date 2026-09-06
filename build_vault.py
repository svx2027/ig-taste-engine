#!/usr/bin/env python3
"""Build content_brain_vault.xlsx from master.csv: an image-embedded, human-editable
review spreadsheet. Cover + up to 3 picked frames per row, all captured/enriched
columns, and yellow-highlighted editable review columns (Liked? / What I liked /
Recreate? / My twist / Priority / Status).

Consolidates four earlier per-run copies of this script into one: the column set
adapts to whatever master.csv actually has (hashtags/mentions, folder_intent,
cross_save_weight, and angle_score/verdict/critique are only added if present),
so the same script works for a single-collection capture or a merged multi-run
master alike.

Usage: python3 build_vault.py [--data-dir DATA_DIR] [--master-csv master.csv]
                               [--out content_brain_vault.xlsx]
"""
import argparse
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

BOX_W, BOX_H = 200, 350  # display box per image (px) — large enough to read overlay text

IMG_COLS = [("thumbnail", "Cover"), ("frame_1", "Frame 1"), ("frame_2", "Frame 2"), ("frame_3", "Frame 3")]

# (key, header, width). Only columns whose key exists in master.csv's header are
# included, except the synthetic ones (num, link) and the editable review
# columns, which are always added.
DATA_COLS_ALL = [
    ("num", "#", 5),
    ("collection", "Collection", 18),
    ("collections", "All collections", 20),
    ("folder_intent", "Intent", 14),
    ("cross_save_weight", "Saves#", 7),
    ("media_type", "Media", 12),
    ("status", "Status", 9),
    ("duration_s", "Len(s)", 7),
    ("title", "Title", 26),
    ("creator_handle", "Creator", 16),
    ("caption", "Caption", 34),
    ("hashtags", "Hashtags", 20),
    ("mentions", "Mentions", 16),
    ("onscreen_text", "On-screen text (verbatim)", 34),
    ("audio_name", "Audio", 22),
    ("lyrics", "Lyrics", 24),
    ("transcript", "Transcript", 26),
    ("likes", "Likes", 9),
    ("comments", "Comments", 9),
    ("shares", "Shares", 7),
    ("post_date", "Posted", 11),
    ("format_template", "Format", 16),
    ("theme", "Theme", 16),
    ("tone", "Tone", 12),
    ("hook_type", "Hook", 16),
    ("audio_type", "Audio type", 14),
    ("language", "Lang", 9),
    ("why_saved", "Why saved", 16),
    ("shareability", "Share-ability", 11),
    ("about", "About", 36),
    ("why_it_works", "Why it works", 32),
    ("recreation_angle", "Recreation angle", 38),
    ("angle_score", "Score", 6),
    ("angle_verdict", "Verdict", 9),
    ("angle_critique", "Angle critique", 34),
]
EDITABLE_COLS = [
    ("liked", "Liked?", 7),
    ("what_i_liked", "What I liked", 26),
    ("recreate", "Recreate?", 10),
    ("my_twist", "My twist", 24),
    ("priority", "Priority", 8),
    ("review_status", "Status2", 13),
]
LINK_COL = ("link", "Link", 8)

WRAP = {
    "title", "caption", "hashtags", "onscreen_text", "lyrics", "transcript",
    "about", "why_it_works", "recreation_angle", "angle_critique",
    "what_i_liked", "my_twist", "audio_name", "collections",
}
EDIT = {k for k, h, w in EDITABLE_COLS}


def embed(ws, tmpdir, path, base, anchor):
    if not path:
        return False
    fp = os.path.join(base, path)
    if not (os.path.exists(fp) and os.path.getsize(fp) > 800):
        return False
    try:
        im = PILImage.open(fp).convert("RGB")
        w0, h0 = im.size
        scale = min(BOX_W / w0, BOX_H / h0)
        dw, dh = max(1, int(w0 * scale)), max(1, int(h0 * scale))
        out = os.path.join(tmpdir, anchor.replace(":", "_") + ".jpg")
        im.save(out, "JPEG", quality=88)
        xim = XLImage(out)
        xim.width, xim.height = dw, dh
        ws.add_image(xim, anchor)
        return True
    except Exception:
        return False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default="data")
    ap.add_argument("--master-csv", default="master.csv")
    ap.add_argument("--out", default="content_brain_vault.xlsx")
    args = ap.parse_args()
    base = args.data_dir

    rows = list(csv.DictReader(open(f"{base}/{args.master_csv}", encoding="utf-8")))
    if not rows:
        raise SystemExit(f"{base}/{args.master_csv} has no rows.")
    present = set(rows[0].keys())

    data_cols = [(k, h, w) for k, h, w in DATA_COLS_ALL if k == "num" or k in present]
    all_cols = [(k, h) for k, h in IMG_COLS] + [(k, h) for k, h, w in data_cols] + \
        [(k, h) for k, h, w in EDITABLE_COLS] + [LINK_COL]

    wb = Workbook()
    ws = wb.active
    ws.title = "Vault"
    HF = PatternFill("solid", start_color="1F3864")
    EDITFILL = PatternFill("solid", start_color="FFF2CC")
    HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    CELL = Font(name="Arial", size=10)
    LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
    thin = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([h for k, h in all_cols])
    img_w_chars = round(BOX_W / 7.0) + 1
    for ci, (k, h) in enumerate(all_cols, 1):
        c = ws.cell(1, ci)
        c.fill = HF
        c.font = HFONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for ci, (k, h) in enumerate(IMG_COLS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = img_w_chars
    offset = len(IMG_COLS)
    for j, (k, h, w) in enumerate(data_cols):
        ws.column_dimensions[get_column_letter(offset + 1 + j)].width = w
    offset += len(data_cols)
    for j, (k, h, w) in enumerate(EDITABLE_COLS):
        ws.column_dimensions[get_column_letter(offset + 1 + j)].width = w

    tmpdir = os.path.join(base, "_vault_tmp")
    os.makedirs(tmpdir, exist_ok=True)

    embedded, failed = 0, 0
    r = 1
    for row in rows:
        r += 1
        media = row.get("media_type", "")
        if row.get("status") not in ("ok", "", None):
            media = f"{media} ({row['status']})"

        vals = {k: row.get(k, "") for k, h, w in data_cols if k != "num"}
        vals["num"] = r - 1
        vals["media_type"] = media
        for k, h, w in EDITABLE_COLS:
            vals[k] = ""

        ci = 0
        for j, (k, h, w) in enumerate(data_cols):
            ci = len(IMG_COLS) + 1 + j
            c = ws.cell(r, ci, vals.get(k, ""))
            c.font = CELL
            c.alignment = Alignment(vertical="top", wrap_text=k in WRAP)
            c.border = BORDER
        for j, (k, h, w) in enumerate(EDITABLE_COLS):
            ci = len(IMG_COLS) + len(data_cols) + 1 + j
            c = ws.cell(r, ci, "")
            c.font = CELL
            c.alignment = Alignment(vertical="top")
            c.fill = EDITFILL
            c.border = BORDER
        link_ci = len(IMG_COLS) + len(data_cols) + len(EDITABLE_COLS) + 1
        lc = ws.cell(r, link_ci, "open")
        lc.hyperlink = row.get("reel_url", "")
        lc.font = LINK
        lc.border = BORDER

        cover = row.get("thumbnail", "") or row.get("frame_1", "")
        imgvals = [cover, row.get("frame_1", ""), row.get("frame_2", ""), row.get("frame_3", "")]
        for ci2, (k, h) in enumerate(IMG_COLS, 1):
            ok = embed(ws, tmpdir, imgvals[ci2 - 1], base, f"{get_column_letter(ci2)}{r}")
            cell = ws.cell(r, ci2)
            cell.border = BORDER
            if ok:
                embedded += 1
            else:
                failed += 1
                cell.value = "" if ci2 > 1 else "(no image)"
                cell.font = Font(name="Arial", size=9, italic=True, color="BBBBBB")
                cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.row_dimensions[r].height = (BOX_H + 8) * 0.75
        if r % 200 == 0:
            print(f"  row {r - 1}/{len(rows)} ...")

    last = ws.max_row
    keyidx = {k: i + 1 for i, (k, h) in enumerate(all_cols)}
    for key, formula in {
        "recreate": '"Yes,Maybe,No"',
        "priority": '"1,2,3"',
        "review_status": '"Idea,Scripting,Filmed,Published"',
        "liked": '"y,n"',
    }.items():
        if key not in keyidx:
            continue
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(keyidx[key])
        dv.add(f"{col}2:{col}{last}")

    ws.freeze_panes = f"{get_column_letter(len(IMG_COLS) + 1)}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(all_cols))}{last}"
    out_path = f"{base}/{args.out}"
    wb.save(out_path)
    print(f"Saved {out_path}")
    print(f"Rows: {last - 1}, Cols: {len(all_cols)}")
    print(f"Images embedded: {embedded}, failed: {failed}")


if __name__ == "__main__":
    main()
